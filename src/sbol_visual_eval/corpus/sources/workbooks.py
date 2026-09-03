"""Parsers for the published overview and yearly annotation workbooks."""

from __future__ import annotations

import hashlib
import re
from typing import Any

from openpyxl import load_workbook

from ..config import MONTH_NUMBER, RETROSPECTIVE_DOI, RUBRIC_URL, YEAR_SUPPLEMENT, YEARS
from ..layout import Layout


def _is_count(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def extract_yearly_records(layout: Layout) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for year in YEARS:
        supplement_number = YEAR_SUPPLEMENT[year]
        filename = f"sb5c00417_si_{supplement_number:03d}.xlsx"
        path = layout.annotations / filename
        if not path.exists():
            raise FileNotFoundError(f"missing {path}; run `sbol-visual-data acquire` first")
        workbook = load_workbook(path, read_only=False, data_only=False)
        for worksheet in workbook.worksheets:
            if worksheet.title == "Summary":
                continue
            if worksheet.title not in MONTH_NUMBER:
                raise ValueError(f"unexpected worksheet {worksheet.title!r} in {filename}")
            for row_number in range(1, worksheet.max_row + 1):
                title = worksheet.cell(row_number, 1).value
                values = [worksheet.cell(row_number, column).value for column in range(3, 7)]
                if not isinstance(title, str) or not all(_is_count(value) for value in values):
                    continue
                total, compatible, compliant, best = (int(value) for value in values)
                comments = []
                for column in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row_number, column)
                    if cell.comment and cell.comment.text.strip():
                        comments.append(
                            {
                                "cell": cell.coordinate,
                                "author": cell.comment.author,
                                "text": cell.comment.text.strip(),
                            }
                        )
                source_key = hashlib.sha256(
                    f"{year}\0{worksheet.title}\0{row_number}\0{title}".encode()
                ).hexdigest()[:16]
                records.append(
                    {
                        "record_id": f"source:{source_key}",
                        "year": year,
                        "issue_month": worksheet.title,
                        "issue_month_number": MONTH_NUMBER[worksheet.title],
                        "title_source": re.sub(r"\s+", " ", title).strip(),
                        "figures_total": total,
                        "figures_sbol_visual_compatible": compatible,
                        "figures_sbol_visual_compliant": compliant,
                        "figures_best_practices": best,
                        "has_compatible_figures": compatible > 0,
                        "has_compliant_figures": compliant > 0,
                        "all_compatible_figures_compliant": compatible > 0
                        and compatible == compliant,
                        "has_best_practice_figures": best > 0,
                        "all_compatible_figures_best_practice": compatible > 0
                        and compatible == compliant == best,
                        "label_granularity": "paper_aggregate_counts",
                        "figure_identities_available": False,
                        "annotation_comments": comments,
                        "annotation_source": f"doi:{RETROSPECTIVE_DOI}.s{supplement_number:03d}",
                        "annotation_path": str(path.relative_to(layout.root)),
                        "annotation_sheet": worksheet.title,
                        "annotation_row": row_number,
                        "rubric_url": RUBRIC_URL,
                    }
                )
    return records


def parse_overview_workbook(layout: Layout) -> dict[int, dict[str, int]]:
    path = layout.annotations / "sb5c00417_si_001.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    output: dict[int, dict[str, int]] = {}
    for row in workbook.active.iter_rows(values_only=True):
        if row and isinstance(row[0], int) and row[0] in YEARS:
            output[row[0]] = {
                "figures_total": int(row[1]),
                "figures_sbol_visual_compatible": int(row[2]),
                "figures_sbol_visual_compliant": int(row[3]),
                "figures_best_practices": int(row[4]),
                "papers_total": int(row[10]),
            }
    return output


def parse_yearly_cached_summaries(layout: Layout) -> dict[int, dict[str, int]]:
    output: dict[int, dict[str, int]] = {}
    for year in YEARS:
        supplement_number = YEAR_SUPPLEMENT[year]
        path = layout.annotations / f"sb5c00417_si_{supplement_number:03d}.xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook["Summary"]
        total_row = next(
            row for row in worksheet.iter_rows(values_only=True) if row[0] == "Total Analyzed"
        )
        output[year] = {
            "figures_total": int(total_row[1]),
            "figures_sbol_visual_compatible": int(total_row[2]),
            "figures_sbol_visual_compliant": int(total_row[3]),
            "figures_best_practices": int(total_row[4]),
            "papers_with_compatible_figures": int(total_row[5]),
            "papers_with_compliant_figures": int(total_row[6]),
            "papers_all_compatible_compliant": int(total_row[7]),
            "papers_with_best_practice_figures": int(total_row[8]),
            "papers_all_compatible_best_practice": int(total_row[9]),
            "papers_total": int(total_row[10]),
        }
    return output
